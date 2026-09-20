from docx import Document
p=r'C:\Users\Windows\Desktop\导诊机器人需求规格说明.docx'
d=Document(p)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))
print('title',d.core_properties.title)
print('headings')
for x in d.paragraphs:
    if x.style.name.startswith('Heading'):
        print('-', x.style.name, x.text[:80])
print('tables rows/cols',[(len(t.rows),len(t.columns)) for t in d.tables[:12]])
text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
for needle in ['基因','健康咨询','用药指导','王小明','110101','急诊医学科','待产品确认','P0-01']:
    print(needle, needle in text)
print('chars', len(text))
