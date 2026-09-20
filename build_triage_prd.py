from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from pathlib import Path

OUT = Path(r'C:\Users\Windows\Desktop\导诊机器人需求规格说明.docx')

# ---------- basic helpers ----------
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ['sz', 'val', 'color', 'space']:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))

def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_col_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(width_inches * 1440)))
    tcW.set(qn('w:type'), 'dxa')

def set_run_font(run, size=10.5, bold=False, color=None, name='Microsoft YaHei'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def style_paragraph(p, space_after=4, line=1.25, first_line=0):
    fmt = p.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Inches(first_line)

def add_text(p, text, size=10.5, bold=False, color=None):
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return r

def add_bullet(doc, text, level=0, style='List Bullet'):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    style_paragraph(p, space_after=2, line=1.2)
    add_text(p, text, size=10.2)
    return p

def add_para(doc, text='', bold_prefix=None, size=10.5, space_after=5, first_line=0):
    p = doc.add_paragraph()
    style_paragraph(p, space_after=space_after, line=1.3, first_line=first_line)
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, size=size, bold=True)
        add_text(p, text[len(bold_prefix):], size=size)
    else:
        add_text(p, text, size=size)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 12.5 if level == 2 else 11, bold=True, color='1F4E79')
    return p

def add_note(doc, label, text, fill='EAF2F8', color='1F4E79'):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c = t.cell(0,0)
    set_col_width(c, 6.75)
    set_cell_shading(c, fill)
    set_cell_margins(c, top=110, start=130, bottom=110, end=130)
    p = c.paragraphs[0]
    style_paragraph(p, space_after=0, line=1.2)
    add_text(p, label + '：', size=10, bold=True, color=color)
    add_text(p, text, size=10, color='333333')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def add_table(doc, headers, rows, widths=None, font_size=8.9, header_fill='D9EAF7', first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = 'Table Grid'
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, h in enumerate(headers):
        cell = hdr.cells[j]
        if widths: set_col_width(cell, widths[j])
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, space_after=0, line=1.1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, h, size=font_size, bold=True, color='1F1F1F')
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            cell = cells[j]
            if widths: set_col_width(cell, widths[j])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if ridx % 2 == 1:
                set_cell_shading(cell, 'F8FBFD')
            p = cell.paragraphs[0]
            style_paragraph(p, space_after=0, line=1.13)
            add_text(p, str(val), size=font_size, bold=(first_col_bold and j == 0))
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={'val':'single','sz':'4','color':'B7C9D6'}, bottom={'val':'single','sz':'4','color':'B7C9D6'}, left={'val':'single','sz':'4','color':'B7C9D6'}, right={'val':'single','sz':'4','color':'B7C9D6'})
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

def add_flow_row(doc, items, fill='EEF5FB'):
    table = doc.add_table(rows=1, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [6.75/len(items)]*len(items)
    for idx, item in enumerate(items):
        c = table.cell(0, idx)
        set_col_width(c, widths[idx])
        set_cell_shading(c, fill)
        set_cell_margins(c, top=120, start=70, bottom=120, end=70)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, space_after=0, line=1.15)
        add_text(p, item, size=9.2, bold=True, color='1F4E79')
        set_cell_border(c, top={'val':'single','sz':'6','color':'9FBAD0'}, bottom={'val':'single','sz':'6','color':'9FBAD0'}, left={'val':'single','sz':'6','color':'9FBAD0'}, right={'val':'single','sz':'6','color':'9FBAD0'})
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

def add_page_field(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)
    set_run_font(run, size=9, color='666666')

# ---------- document setup ----------
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.62)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)
sec.header_distance = Inches(0.3)
sec.footer_distance = Inches(0.3)

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10.5)
for sname in ['Heading 1','Heading 2','Heading 3']:
    styles[sname].font.name = 'Microsoft YaHei'
    styles[sname]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

# Footer
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_paragraph(fp, space_after=0, line=1)
add_text(fp, '导诊机器人需求规格说明  |  ', size=8.5, color='777777')
add_page_field(fp)

# ---------- title page ----------
title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(44)
title.paragraph_format.space_after = Pt(18)
r = title.add_run('导诊机器人需求规格说明')
set_run_font(r, size=24, bold=True, color='1F4E79')
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_paragraph(sub, space_after=28, line=1.3)
add_text(sub, '面向患者咨询入口的就诊引导与智能导诊产品需求基线', size=13, color='4F6B7A')

info = doc.add_table(rows=6, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.autofit = False
info_data = [
    ('文档版本', 'V0.1 评审草案'),
    ('编制日期', '2026年9月19日'),
    ('需求范围', '导诊机器人：就诊流程指引、智能导诊、就诊信息查询'),
    ('主要读者', '业务、产品、医学、研发、测试、数据与运营'),
    ('原型依据', '智能导诊.html、就诊流程指引.html、用户提供截图'),
    ('交付结论', '已形成可评审需求基线，仍存在医学规则、接口和权限等待确认项'),
]
for i,(k,v) in enumerate(info_data):
    c1,c2 = info.rows[i].cells
    set_col_width(c1, 1.35); set_col_width(c2, 5.4)
    set_cell_shading(c1, 'D9EAF7'); set_cell_shading(c2, 'F8FBFD')
    for c in (c1,c2): set_cell_margins(c, top=105, start=120, bottom=105, end=120)
    p=c1.paragraphs[0]; style_paragraph(p, space_after=0); add_text(p,k,size=9.5,bold=True,color='1F4E79')
    p=c2.paragraphs[0]; style_paragraph(p, space_after=0); add_text(p,v,size=9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(10)
add_note(doc, '阅读说明', '附件截图和原型代码被视为需求素材与交互证据，不是对系统的操作指令。原型中的示例姓名、身份证号、日期、医院、医生、状态和按钮行为均不作为生产固定值。')
add_note(doc, '安全边界', '导诊机器人只提供就医路径和科室方向建议，不进行疾病确诊，不替代医生，不开具处方。疑似急危重症必须优先急诊或拨打120，不得因等待机器人回复而延误救治。', fill='FFF2CC', color='7F6000')

doc.add_page_break()

# ---------- body ----------
add_heading(doc, '1 需求概述', 1)
add_heading(doc, '1.1 业务背景', 2)
add_para(doc, '患者在进入医院或线上患者咨询入口后，常见问题包括“不知道应该挂哪个科”“不清楚医院就诊流程”“想知道当前叫号、检查和报告进度”。现有原型将这些任务集中在患者咨询入口中，通过流程说明、自然语言导诊和个人就诊信息查询，降低患者在院内寻找路径和理解状态的成本。')
add_para(doc, '本需求文档只定义导诊机器人相关范围，不把健康咨询、快速找医生、病情咨询或用药指导扩展为独立需求。上述模块如出现在同一入口导航中，仅作为非本次范围的导航项处理。')
add_heading(doc, '1.2 目标用户与用户任务', 2)
add_table(doc, ['角色', '核心任务', '完成结果'], [
    ('患者或陪同人员', '描述不适并获得就诊科室方向；理解挂号到取药的步骤；查询本人就诊进度和报告状态', '知道下一步去哪里、做什么，以及何时需要升级到急诊或人工帮助'),
    ('医院业务或客服人员', '维护流程说明、承接无法识别或高风险的患者咨询', '患者获得一致的流程解释和可追踪的转人工入口'),
    ('医学审核人员', '审核高风险触发词、导诊方向和安全提示', '导诊规则可审阅、可版本化，且不把原型演示逻辑直接当作医学结论'),
    ('研发与测试人员', '按字段、状态、异常和验收规则实现与验证', '能够构建端到端流程、接口异常和安全边界用例'),
], widths=[1.25,3.0,2.5], font_size=8.8, first_col_bold=True)
add_heading(doc, '1.3 产品目标与成功标准', 2)
for x in [
    '患者能够从入口明确选择“看流程”“做导诊”或“查就诊信息”，不需要从多个页面反向猜测使用方式。',
    '智能导诊能够基于患者主诉进行有限轮次追问，输出建议就诊科室或明确的急诊升级提示，并持续展示安全边界。',
    '就诊流程指引能够覆盖挂号、候诊、检查、缴费、取药五个阶段，说明每一步的目的、患者操作和注意事项。',
    '就诊信息查询能够在身份核验和医院接口成功的前提下展示就诊记录、当前状态、叫号和报告状态，并在失败时给出可理解的恢复路径。',
    '所有原型示例数据在生产环境中被替换为真实接口数据、配置数据或空态，不会以演示值混入患者数据。',
]: add_bullet(doc, x)

add_heading(doc, '1.4 信息可信度与需求标记', 2)
add_table(doc, ['标记', '本文件中的含义', '示例'], [
    ('已确认事实', '在用户指令、截图或原型代码中明确存在，可作为当前需求基线', '存在智能导诊页面；原型包含六个导诊方向；查询表单包含姓名和身份证号码'),
    ('工作假设', '为了形成可评审方案暂时采用，需在设计或研发前验证', '首期为患者侧网页或移动端入口；导诊结果默认只推荐科室，不直接挂号'),
    ('建议', '产品优化或风险控制建议，不等同于已确认规则', '高风险命中后提供急诊、120和人工帮助入口；不要把自然语言误匹配强制映射为首个选项'),
    ('待产品确认', '缺失信息会影响范围、接口、权限、医学规则或验收，必须在相应门禁前确认', '急诊规则责任人、医院接口、身份核验方式、数据保存周期'),
], widths=[1.1,2.55,3.1], font_size=8.6, first_col_bold=True)

add_heading(doc, '2 范围与约束', 1)
add_heading(doc, '2.1 本次范围', 2)
add_table(doc, ['范围项', '包含内容', '优先级'], [
    ('就诊流程指引', '展示挂号、候诊、检查、缴费、取药的步骤说明，以及患者操作和注意事项；允许跳转到智能导诊或人工咨询入口（若配置）', 'P0'),
    ('智能导诊', '症状文本输入、多轮追问、快速回复、风险识别、科室方向建议、急诊升级、重新导诊和补充描述', 'P0'),
    ('就诊信息查询', '姓名和证件信息输入、身份核验、就诊记录列表、进行中就诊详情、叫号、检查/报告状态、查询异常处理', 'P1，依赖医院接口'),
    ('统一安全与审计', '免责声明、敏感字段脱敏、访问日志、规则和模型版本留痕、异常监控', 'P0'),
], widths=[1.45,4.55,0.75], font_size=8.7, first_col_bold=True)
add_heading(doc, '2.2 非本次范围', 2)
for x in [
    '疾病诊断、检验结果解读、治疗方案推荐、处方开具和用药指导。',
    '健康咨询、病情咨询、快速找医生等其他入口模块的完整需求；这些模块如存在，仅保留导航关系。',
    '直接挂号、缴费、预约检查、在线问诊或在线处方，除非后续明确新增范围并补充接口与责任边界。',
    '医院内部排班、叫号机、HIS、LIS、RIS、PACS、支付和药房系统的内部改造；本文件只定义导诊机器人对这些系统的业务依赖和展示要求。',
]: add_bullet(doc, x)
add_heading(doc, '2.3 依赖、约束与主要风险', 2)
add_table(doc, ['类别', '内容', '影响'], [
    ('医学依赖', '急诊高风险词、分诊方向、提示文案和漏诊/误报策略需由医学专家审核', '未确认前不能将原型关键词作为生产医学规则'),
    ('系统依赖', '就诊查询依赖医院挂号、叫号、检查、报告和药房等系统接口', '接口不可用时只能展示错误或最近一次允许展示的数据，不能伪造状态'),
    ('隐私约束', '姓名、身份证号、就诊历史、报告和处方信息属于敏感个人信息', '必须进行身份核验、最小化展示、脱敏、审计和访问控制'),
    ('产品责任', '导诊建议可能被患者理解为诊断或唯一就医路径', '必须明确“建议而非诊断”，高风险场景不得延误急救'),
    ('原型限制', '当前原型包含模拟延时、关键词分类、默认演示数据和固定状态', '原型行为不能直接等同于生产算法、性能目标或接口协议'),
], widths=[1.1,3.4,2.25], font_size=8.6, first_col_bold=True)

add_heading(doc, '3 页面地图与端到端流程', 1)
add_heading(doc, '3.1 页面地图', 2)
add_table(doc, ['入口或页面', '主要作用', '与本需求关系'], [
    ('患者咨询入口', '承载患者进入流程指引、智能导诊和就诊信息查询的导航', '本需求的统一上游入口'),
    ('就诊流程指引', '解释从挂号到取药的通用流程', '本需求范围内页面'),
    ('智能导诊', '通过主诉和追问推荐就诊科室或急诊', '本需求核心页面'),
    ('就诊信息查询', '验证本人身份并展示就诊进度、叫号和报告状态', '本需求范围内页面，依赖医院接口'),
    ('人工帮助或线下窗口', '承接机器人无法识别、接口失败、疑似急危重症或患者需要人工帮助的场景', '外部交接对象，具体入口待确认'),
], widths=[1.45,3.1,2.2], font_size=8.7, first_col_bold=True)
add_heading(doc, '3.2 总体业务流程', 2)
add_para(doc, '患者从患者咨询入口进入后，先选择任务类型。流程指引是静态解释路径；智能导诊是对话和风险分流路径；就诊信息查询是身份核验和医院数据展示路径。三条路径均应在无法继续时给出清晰的人工或线下就医建议，不能以空白页面结束。')
add_flow_row(doc, ['进入患者咨询入口', '选择任务', '流程指引 / 智能导诊 / 信息查询', '系统校验与处理', '展示结果或安全升级', '结束或转人工/线下'])
add_flow_row(doc, ['描述流程需求', '输入症状主诉', '身份信息查询'], fill='F7F7F7')
add_para(doc, '智能导诊主路径：输入主诉 → 识别可用导诊方向 → 追问关键事实 → 判断是否命中高风险 → 输出科室方向或急诊建议 → 患者可补充描述或重新导诊。就诊信息查询主路径：输入姓名和证件号码 → 字段校验 → 身份核验 → 拉取就诊记录 → 查看当前就诊详情和报告状态 → 刷新或结束。', size=10, space_after=6)
add_heading(doc, '3.3 状态模型总览', 2)
add_table(doc, ['业务对象', '主要状态', '进入条件', '允许操作或结果'], [
    ('导诊会话', '未开始、追问中、分析中、已完成、无法识别、已重置', '进入页面、提交主诉、完成追问、分析返回或用户点击重置', '输入/选择回复、等待分析、查看建议、继续补充、重新开始'),
    ('就诊信息查询', '未查询、校验失败、查询中、查询成功、无记录、接口失败', '用户填写并提交身份信息，或接口返回结果', '修改字段、重试、查看记录、查看详情；失败不可展示伪造数据'),
    ('就诊记录', '进行中、已完成、状态未知', '由医院侧就诊数据提供', '进行中可查看实时状态；已完成可回顾历史；未知需提示更新时间或联系医院'),
    ('报告项目', '未开立、待检查、检查中、待生成、已生成、不可用', '由医院检查和报告系统提供', '查看状态；只有已生成且有权限时允许查看报告详情'),
], widths=[1.25,1.45,2.25,1.8], font_size=8.35, first_col_bold=True)

# 4 process guide
add_heading(doc, '4 功能需求', 1)
add_heading(doc, '4.1 就诊流程指引', 2)
add_para(doc, '该页面用于帮助不熟悉医院流程的患者先理解“先做什么、再做什么”。页面不应让患者误以为所有医院都完全按照同一流程运行，因此内容需标识为通用指引，并允许医院按院区、门诊类型或配置版本调整。')
add_heading(doc, '4.1.1 页面结构与展示规则', 3)
add_table(doc, ['需求编号', '规则', '来源与备注'], [
    ('GUIDE-001', '页面展示五个主步骤：挂号、候诊、检查、缴费、取药；每步至少包含步骤名称、这一步做什么、患者需要做什么、注意事项。', '截图和就诊流程指引原型已确认'),
    ('GUIDE-002', '挂号步骤说明选择科室、医生、时间和就诊顺序；不确定科室时提示先使用智能导诊或咨询台确认。', '原型事实；直接挂号不在本期范围'),
    ('GUIDE-003', '候诊步骤说明到指定楼层、诊区或诊室附近等待，保持手机畅通并关注叫号；错过叫号时提示咨询护士台。', '原型事实'),
    ('GUIDE-004', '检查步骤至少覆盖抽血、影像、心电图、超声等示例类型，并提示预约、缴费、空腹、憋尿等要求以医院现场规则为准。', '原型事实；示例类型不可视为固定清单'),
    ('GUIDE-005', '缴费步骤提示患者核对姓名、项目和金额，并按医院支持的窗口、自助机、手机端或医保方式支付。', '原型事实；支付接口不在本期范围'),
    ('GUIDE-006', '取药步骤提示核对姓名、药品名称、数量和用法，疑问时向药师咨询；不得在本页面给出个体化用药建议。', '安全边界'),
    ('GUIDE-007', '流程内容应支持版本化或配置化；展示内容必须标识适用医院/院区和更新时间（若后台提供）。', '建议，待后台能力确认'),
], widths=[1.0,4.25,1.5], font_size=8.35, first_col_bold=True)
add_heading(doc, '4.1.2 页面状态与异常', 3)
add_table(doc, ['状态', '触发条件', '系统表现', '恢复方式'], [
    ('加载中', '页面或配置正在加载', '展示加载提示，不展示不完整步骤', '加载成功自动展示；超时显示重试'),
    ('正常', '内容加载成功', '展示五步流程和注意事项', '患者可切换步骤或跳转智能导诊'),
    ('内容为空', '医院配置未发布或无可用版本', '展示通用流程兜底或明确“暂未配置”，不得出现空白', '重试或联系医院管理员'),
    ('加载失败', '网络、配置服务或接口异常', '提示暂时无法加载，并提供重试和线下咨询建议', '用户点击重试；重复失败时转人工/服务台'),
], widths=[1.05,2.05,2.55,1.1], font_size=8.5, first_col_bold=True)

# 4.2 triage
add_heading(doc, '4.2 智能导诊', 2)
add_para(doc, '智能导诊是本产品的核心功能。患者用自然语言描述主要不适，系统通过受控的多轮问答补充部位、持续时间、伴随表现和风险信号，最终输出“建议就诊科室”或“立即急诊/拨打120”的路径建议。系统必须把导诊结果定义为就医路径建议，而不是诊断结果。')
add_heading(doc, '4.2.1 页面组成与输入', 3)
add_table(doc, ['元素', '规则', '校验与反馈'], [
    ('安全提示', '页面固定展示胸痛、呼吸困难、大出血、意识模糊、严重外伤等高风险示例，并提示急诊/120优先。', '文案由医学审核；高风险提示不得被用户关闭后永久隐藏'),
    ('主诉输入框', '支持患者手动输入主要不适，提示部位、持续时间、严重程度等信息。', '去除首尾空格；原型中少于2个字符时提示补充；空输入不可提交'),
    ('发送按钮', '提交主诉或当前轮次的手动回复。', '分析中禁用；提交后清空输入或按产品确认保留草稿'),
    ('快速回复', '系统根据当前问题展示可点击选项；点击后作为用户回复进入会话。', '点击后立即进入下一轮；避免重复提交'),
    ('对话区', '按时间顺序展示用户消息、机器人问题、分析中提示、结果和错误提示。', '区分消息角色和状态；长文本可换行，不截断关键信息'),
    ('重新导诊', '清空当前对话和当前会话状态，回到初始欢迎语。', '二次确认是否需要待确认；重置后不能继续使用旧会话的结果'),
], widths=[1.2,3.75,1.8], font_size=8.35, first_col_bold=True)
add_heading(doc, '4.2.2 原型已表达的导诊方向', 3)
add_note(doc, '原型证据边界', '原型当前通过关键词进入六个演示剧本：急诊医学科、呼吸内科、神经内科、消化内科、骨科、皮肤科。以下内容只能作为首期评审样例和交互覆盖范围，不能直接视为最终医学分诊规则、完整科室清单或固定关键词表。', fill='F3F6F8', color='4F6B7A')
add_table(doc, ['演示方向', '原型典型主诉', '原型追问/风险信号', '生产要求'], [
    ('急诊医学科', '胸痛、大汗、压榨感、撕裂感、窒息、昏迷、抽搐等', '放射痛、呼吸困难、意识异常等高危表现', '由医学专家定义急诊升级条件和优先级'),
    ('呼吸内科', '咳嗽、咳痰、胸闷、气短、发热、咽痛等', '夜间或活动后加重；明显呼吸困难、持续胸痛、口唇发绀时升级', '确认科室映射和儿童/老年等特殊人群规则'),
    ('神经内科', '头痛、头晕、失眠、焦虑、麻木等', '头痛性质、头晕、恶心；突发剧烈头痛、肢体无力、意识异常时升级', '确认神经急症的医学规则和人工转接'),
    ('消化内科', '腹痛、腹泻、恶心、呕吐、胃胀、反酸等', '部位、绞痛/隐痛；持续剧烈腹痛、便血、明显脱水时升级', '确认急腹症和儿童患者规则'),
    ('骨科', '关节痛、腰痛、扭伤、骨折等', '外伤、肿胀、活动受限；畸形或无法活动时升级', '确认创伤急救和院前建议边界'),
    ('皮肤科', '皮疹、瘙痒、红疹等', '位置、红肿、扩散；口唇肿胀、呼吸困难、快速扩散时升级', '确认过敏反应和急诊规则'),
], widths=[1.0,2.0,1.8,2.0], font_size=7.9, first_col_bold=True)
add_heading(doc, '4.2.3 多轮对话与状态规则', 3)
add_table(doc, ['需求编号', '前置条件', '系统行为', '完成或异常结果'], [
    ('TRIAGE-001', '首次提交主诉', '提取症状线索并选择候选导诊剧本；当前原型为关键词分类，生产算法和词表需医学审核。', '识别成功进入追问；识别失败提示重新描述，不得默认进入任意科室'),
    ('TRIAGE-002', '进入追问', '一次展示一个问题，提供快速回复；同时允许患者手动输入。', '记录症状、持续时间、伴随表现和风险信号'),
    ('TRIAGE-003', '用户点击快速回复', '将选项作为当前轮次答案，进入下一轮或分析。', '防止重复点击；分析中禁用回复控件'),
    ('TRIAGE-004', '用户手动输入追问答案', '按医学审核后的语义匹配规则处理；无法匹配时应请求澄清。', '不得沿用原型“无法匹配则取第一个选项”的演示行为，除非产品和医学明确接受'),
    ('TRIAGE-005', '高风险信号命中', '停止普通科室导诊或将急诊升级置于结果首位，明确立即前往急诊或拨打120。', '高风险提示必须可见；不得要求患者继续完成非必要问答'),
    ('TRIAGE-006', '完成分析', '展示建议科室、依据摘要、风险提示、免责声明和下一步行动。', '结果状态为已完成；是否允许直接挂号待确认'),
    ('TRIAGE-007', '完成后继续补充', '原型允许继续输入但不重新计算旧结论。生产版本应提示“补充信息是否重新评估”。', '待确认；至少不得静默覆盖原结果'),
    ('TRIAGE-008', '点击重新导诊', '清空消息、追问轮次、摘要、风险状态和临时输入，回到欢迎语。', '新会话与旧结果隔离；旧结果是否保存由数据策略确认'),
], widths=[1.0,1.6,3.1,1.1], font_size=8.0, first_col_bold=True)
add_heading(doc, '4.2.4 导诊结果展示', 3)
add_table(doc, ['结果类型', '必须展示', '不得展示或承诺'], [
    ('建议科室', '建议就诊科室；已识别症状、持续时间和伴随表现摘要；下一步就医提示；必要的复诊或急诊升级提醒；安全免责声明', '不得写成“确诊”“一定是某病”；不得给出具体处方、剂量或治疗方案'),
    ('急诊升级', '醒目风险提示；立即前往急诊医学科或拨打120；不要等待、不要自行前往其他专科；如需要可提供人工帮助', '不得继续强制追问；不得等待模型置信度达到某数值再提示'),
    ('无法识别', '说明未识别到足够有效信息；引导患者补充部位、时间、严重程度和伴随表现；提供人工或线下咨询入口（若配置）', '不得随机推荐一个科室或把示例内容当作患者事实'),
    ('多科室可能', '说明存在多个可能的就医方向，提出最少必要澄清问题或建议综合/人工咨询', '不得隐瞒不确定性或展示伪造的唯一结论'),
], widths=[1.05,4.0,1.75], font_size=8.4, first_col_bold=True)
add_heading(doc, '4.2.5 智能导诊异常处理', 3)
add_table(doc, ['异常', '触发条件', '系统反馈', '恢复方式'], [
    ('空输入', '输入为空或仅包含空白字符', '不发送消息，提示输入主要不适', '补充后重试'),
    ('内容过短', '原型规则为去空格后少于2个字符', '提示补充症状、部位或表现', '补充后重试；长度上限待确认'),
    ('无法识别', '未命中有效方向或信息不足', '提示用自然语言说明部位、持续时间和表现，并提供人工/线下入口', '重新输入或转人工'),
    ('模型/规则服务超时', '分析服务超过约定时间', '显示分析失败，不展示旧结果或模拟结果', '重试；多次失败转人工'),
    ('用户重复提交', '分析中点击发送或快速回复', '按钮禁用，保留一次请求', '分析完成或失败后恢复'),
    ('会话过期', '页面长时间未操作或服务端会话过期', '提示会话已过期，是否重新开始', '重新导诊；旧会话不再继续写入'),
], widths=[1.15,1.9,2.6,1.1], font_size=8.25, first_col_bold=True)

# 4.3 query
add_heading(doc, '4.3 就诊信息查询', 2)
add_para(doc, '该模块用于医院系统完成对接后，让患者查询本人的就诊记录、进行中状态、叫号和报告状态。截图中的姓名“王小明”、证件号“110101********1234”、医院名称、医生、日期和诊断均为演示值，生产环境必须来自经过授权的医院侧数据。')
add_heading(doc, '4.3.1 输入与身份核验', 3)
add_table(doc, ['字段或操作', '类型与要求', '校验、展示和权限规则'], [
    ('患者姓名', '必填文本', '去除首尾空格；禁止空值提交；按医院接口要求进行规范化；结果页仅展示与当前身份核验一致的患者姓名'),
    ('身份证号码或证件号', '必填敏感字段', '输入格式、长度和证件类型由医院接口确认；页面输入可按需要遮挡；结果页和日志不得明文暴露完整号码'),
    ('查询就诊信息', '按钮', '仅在必填项通过基础校验后可提交；查询中禁用，避免重复请求；不提供未授权的模糊查询'),
    ('就诊记录列表', '接口返回集合', '展示就诊日期、医院/院区、科室、医生、当前状态；演示固定记录不能作为生产固定条数'),
    ('查看详情', '记录级操作', '只能查看通过当前身份核验且有权限的记录；进行中记录显示实时状态，已结束记录显示可回顾内容'),
], widths=[1.35,1.9,3.5], font_size=8.35, first_col_bold=True)
add_heading(doc, '4.3.2 查询结果与状态展示', 3)
add_table(doc, ['信息区域', '展示内容', '数据来源与规则'], [
    ('就诊信息概览', '姓名、脱敏证件号、医院/院区、科室、医生、就诊时间', '医院就诊接口；按最小必要原则展示，历史数据是否全部返回待确认'),
    ('当前进度', '挂号、候诊、医生接诊、检查、缴费、取药等阶段的当前状态和下一步提示', '挂号、叫号、检查、支付、药房接口；若状态口径不一致需建立映射表'),
    ('叫号信息', '当前叫号、患者排队号、诊室/检查室/药房窗口等', '叫号系统；数据过期时显示更新时间，不能保证实时则明确提示'),
    ('报告状态', '未开立、待检查、检查中、待生成、已生成、不可用等', '检验检查和报告系统；只有授权且已生成时才允许查看详情'),
    ('就诊小结', '诊断、处方、医嘱、补充说明等', '医生或医院系统；字段未生成时展示“待医生填写”，不得由机器人生成医疗结论'),
], widths=[1.25,2.7,2.8], font_size=8.25, first_col_bold=True)
add_heading(doc, '4.3.3 查询状态与异常', 3)
add_table(doc, ['状态或异常', '系统行为', '用户反馈与恢复'], [
    ('未查询', '展示姓名和证件输入区域及查询说明', '用户补全信息后提交'),
    ('字段校验失败', '不调用医院接口', '标记具体错误字段，用户修改后重试'),
    ('查询中', '显示进度或加载状态，锁定重复提交', '接口返回、超时或失败后恢复操作'),
    ('身份不匹配', '不返回任何就诊记录详情', '提示姓名与证件信息不匹配，建议核对或联系医院'),
    ('查无记录', '不展示空的详情卡片或演示数据', '提示当前没有可查询记录，可修改查询条件或咨询医院'),
    ('接口不可用/超时', '不伪造状态，不用旧缓存冒充实时结果', '说明暂时无法查询，允许重试，并提供线下窗口/人工咨询建议'),
    ('部分接口失败', '展示已成功获得的数据并标注缺失区域，或按安全策略整体失败', '明确哪些数据未更新和最后更新时间；恢复策略待接口合同确认'),
    ('无权限或授权失效', '拒绝访问敏感详情', '提示重新完成身份核验或联系医院，不显示接口内部错误'),
], widths=[1.4,3.0,2.35], font_size=8.2, first_col_bold=True)

# 5 cross-page
add_heading(doc, '5 跨页面交互与数据规则', 1)
add_heading(doc, '5.1 页面间传递关系', 2)
add_table(doc, ['来源页面', '目标页面或动作', '传递数据', '规则'], [
    ('就诊流程指引', '进入智能导诊', '入口来源标识，不默认带入个人健康信息', '患者主动触发；进入后从欢迎语开始'),
    ('智能导诊', '返回流程指引或人工帮助', '建议科室、风险等级、会话标识（如有）', '只在患者授权和产品确认后传递；不得把导诊结果当作诊断'),
    ('智能导诊', '挂号或预约入口', '建议科室（可选）', '当前非默认范围，是否直达挂号为阻塞待确认项'),
    ('就诊信息查询', '流程步骤详情', '当前就诊阶段', '可用于定位到挂号、候诊、检查、缴费或取药说明；需防止状态口径不一致'),
    ('任一页面', '刷新/返回', '页面状态和必要的会话标识', '返回不应泄露他人数据；会话过期需重新校验'),
], widths=[1.45,1.7,2.0,1.6], font_size=8.25, first_col_bold=True)
add_heading(doc, '5.2 数据最小化与存储', 2)
for x in [
    '导诊会话至少包含会话标识、用户输入、系统追问、用户回答、结构化摘要、风险命中结果、导诊建议、规则/模型版本和时间戳。是否保存原始主诉、保存多久以及是否允许医生查看，需由隐私和业务确认。',
    '就诊查询只在身份核验成功后展示本人可访问的数据；默认不在前端长期保存完整身份证号码，不在日志中记录完整证件号码。',
    '原型中的示例患者、医院、医生、日期和报告内容必须在生产部署前全部替换为接口数据或配置数据。测试数据应使用脱敏或专门构造的假数据。',
    '任何由医院接口返回的状态都应带有来源、更新时间和必要的状态映射；当多个系统状态冲突时，按接口责任和数据新鲜度定义优先级，并记录异常。',
]: add_bullet(doc, x)

# 6 NFR
add_heading(doc, '6 通用业务规则与非功能要求', 1)
add_heading(doc, '6.1 安全与医学责任', 2)
add_table(doc, ['规则编号', '要求'], [
    ('SAFE-001', '所有导诊结果页面固定展示“本结果仅用于就医路径建议，不构成医学诊断或治疗建议；如症状严重或快速加重，请立即线下就医”的安全说明。'),
    ('SAFE-002', '高风险命中优先级高于普通科室推荐；如果同一输入同时命中普通方向和急诊风险，应以急诊升级为首要结果。'),
    ('SAFE-003', '当模型、规则或接口无法判断时，系统必须明确不确定性，建议人工咨询或线下就医，不得编造科室、状态、报告或诊断。'),
    ('SAFE-004', '急诊提示的触发词、组合条件、敏感度、误报处理、特殊人群和多语言表达由医学负责人审核并版本化。'),
    ('SAFE-005', '系统不得因为用户未登录、接口超时或导诊服务不可用而阻碍患者拨打120或前往急诊。'),
], widths=[1.1,5.65], font_size=8.7, first_col_bold=True)
add_heading(doc, '6.2 隐私、权限与审计', 2)
add_table(doc, ['方面', '要求'], [
    ('功能权限', '患者可使用流程指引和智能导诊；就诊查询需完成身份核验；运营或医学人员的规则维护、日志查看和报告查看权限分离。'),
    ('数据权限', '患者仅能查看本人或已授权代办人的就诊数据；代办人、儿童、老人和家庭成员关系的授权方式待确认。'),
    ('字段权限', '身份证号默认脱敏；报告、诊断、处方和医嘱按最小必要原则展示；导诊会话原文是否可被医生查看需单独授权。'),
    ('审计', '记录查询主体、时间、访问对象、结果状态、异常、规则/模型版本和人工接管事件；敏感字段日志脱敏。'),
    ('留存与删除', '导诊会话和查询日志留存期限、用户删除请求、监管留痕和备份策略待隐私负责人确认。'),
], widths=[1.25,5.5], font_size=8.65, first_col_bold=True)
add_heading(doc, '6.3 性能、可用性与可观测性', 2)
add_table(doc, ['方面', '初步要求', '待确认项'], [
    ('响应体验', '页面交互应给出加载、分析中和失败状态；禁止无反馈等待。', '分析接口、医院接口的P95/P99响应目标'),
    ('可用性', '流程指引在接口不可用时尽量提供已审核的通用兜底内容；就诊查询不得用旧数据伪装实时数据。', '服务等级、降级策略、缓存有效期'),
    ('兼容性', '覆盖项目确定的主流浏览器和移动端；文本换行、键盘输入和辅助技术可用。', '浏览器版本、无障碍等级、语音/图片输入范围'),
    ('监控', '监控空输入、无法识别、急诊命中、接口失败、重复提交、人工转接和异常状态分布。', '告警阈值、脱敏方案、报表口径'),
    ('版本治理', '每次医学规则、文案、模型、状态映射或接口字段变更均可追溯。', '审批人、灰度策略、回滚方案'),
], widths=[1.15,3.25,2.35], font_size=8.35, first_col_bold=True)

# 7 acceptance
add_heading(doc, '7 端到端验收要点', 1)
add_para(doc, '以下验收用例用于覆盖当前范围的正常、异常、权限和安全路径。具体医学触发词和接口返回码在确认后补充到测试数据集和接口合同。')
add_table(doc, ['编号', '前置条件', '操作', '预期结果'], [
    ('AT-001', '流程指引内容可用', '进入页面并依次查看五个步骤', '展示挂号、候诊、检查、缴费、取药；每一步含目的、患者操作和注意事项'),
    ('AT-002', '流程配置加载失败', '刷新页面或模拟配置服务失败', '展示明确失败提示和重试/线下咨询，不出现空白或半截内容'),
    ('AT-003', '智能导诊初始状态', '进入智能导诊页面', '展示安全提示、欢迎语、输入框、发送按钮和重新导诊入口；发送按钮在空输入时不可用'),
    ('AT-004', '用户输入“头痛两天、伴恶心”', '提交主诉并完成追问', '进入神经内科演示方向的追问并输出建议；结果标记为导诊建议而非诊断；具体文案以医学确认版为准'),
    ('AT-005', '用户输入高风险胸痛描述', '提交“胸痛伴大汗、左臂放射痛”等高风险信息', '急诊升级提示置于首位，建议立即前往急诊或拨打120，不要求继续完成普通科室追问'),
    ('AT-006', '用户输入空白或少于2个有效字符', '点击发送或按回车', '不调用分析服务，提示补充主要症状；原型行为可作为回归用例'),
    ('AT-007', '用户输入无法识别的内容', '提交与症状无关或信息过少的描述', '提示重新描述部位、时间和表现，不能随机推荐科室'),
    ('AT-008', '智能导诊正在分析', '连续点击发送和快速回复', '控件禁用或只接受一次请求，不产生重复消息或重复分析'),
    ('AT-009', '导诊结果已完成', '输入补充症状并点击重新导诊', '补充信息按产品确认策略处理；重新导诊后会话清空且旧结论不再冒充当前结果'),
    ('AT-010', '查询表单为空', '点击查询', '不请求医院接口，标记姓名和证件必填错误'),
    ('AT-011', '身份信息不匹配', '输入一组不匹配的姓名和证件号', '不展示任何就诊记录详情，提示核对或联系医院'),
    ('AT-012', '查询成功且有进行中记录', '完成核验并打开一条进行中就诊记录', '展示医院、科室、医生、就诊时间、当前阶段、叫号、报告状态和更新时间；敏感字段脱敏'),
    ('AT-013', '查询成功但无记录', '完成核验后接口返回空集合', '展示无记录空态，不展示原型示例数据，并提供修改条件或人工咨询建议'),
    ('AT-014', '医院接口超时或部分失败', '模拟挂号/叫号/报告接口不可用', '展示可理解错误和重试；不伪造状态，并标明未更新的数据区域'),
    ('AT-015', '未授权用户访问历史详情', '尝试通过旧链接或修改参数访问他人记录', '拒绝访问并记录审计事件，不返回敏感信息'),
    ('AT-016', '生产数据部署前检查', '检索页面、配置、日志和测试环境数据', '不存在示例姓名、完整证件号、固定演示报告或固定结果作为生产数据'),
], widths=[0.7,1.8,2.2,2.05], font_size=7.8, first_col_bold=True)

# 8 open questions
add_heading(doc, '8 待产品确认项与上线门禁', 1)
add_heading(doc, '8.1 阻塞项 P0', 2)
add_table(doc, ['编号', '必须确认的问题', '影响范围', '最迟确认节点'], [
    ('P0-01', '导诊结果是仅推荐科室，还是可以直接生成挂号/预约入口？', '页面流程、接口、责任边界、验收', '产品方案评审前'),
    ('P0-02', '首期支持哪些医院、院区、科室、医生和特殊人群？', '科室映射、数据权限、内容配置', '范围冻结前'),
    ('P0-03', '急诊高风险规则由谁审核，采用哪些组合条件和升级策略？', '医学安全、模型评估、测试集', '医学评审和上线前'),
    ('P0-04', '用户可自由输入、结构化问答，还是两者并存？手动答案如何做语义匹配？', '对话引擎、交互、误匹配风险', '技术方案评审前'),
    ('P0-05', '无法识别、多科室可能和信息冲突时，是否转人工？人工入口和责任人是谁？', '兜底流程、客服系统、运营', '流程评审前'),
    ('P0-06', '就诊信息查询接口是否已存在？接口所有者、鉴权、限流、字段和错误码是什么？', '查询功能能否实现、数据模型和测试', '接口设计前'),
    ('P0-07', '查询身份核验采用姓名+身份证号，还是需要短信/医保/账号等二次认证？', '隐私、越权风险、登录流程', '安全评审前'),
    ('P0-08', '叫号、报告、检查、缴费、取药状态分别来自哪些系统，状态刷新频率和冲突处理如何定义？', '状态展示、数据新鲜度、异常处理', '接口合同前'),
    ('P0-09', '是否允许展示历史就诊、诊断、处方、医嘱和报告详情？患者授权及代办人规则是什么？', '字段权限、合规、页面结构', '权限设计前'),
    ('P0-10', '导诊会话是否保存，保存多久，是否允许医生或客服查看？', '数据留存、审计、隐私', '数据方案评审前'),
    ('P0-11', '医院责任、免责声明、人工接管和急诊责任边界由谁签字确认？', '上线责任和风险控制', '上线门禁前'),
], widths=[0.75,3.0,1.8,1.2], font_size=7.8, first_col_bold=True)
add_heading(doc, '8.2 一般待确认项 P1', 2)
for x in [
    '是否支持语音、图片、方言或多语言输入；当前原型仅证明文本和快速回复。',
    '是否支持会话恢复、跨设备继续、导诊结果分享或导出；如支持，需补充授权和有效期。',
    '对话和查询页面的浏览器、移动端、无障碍、弱网和低端设备支持范围。',
    '查询结果按时间筛选、分页、刷新频率和报告详情的展示范围。',
    '运营后台是否允许配置流程文案、科室映射、风险规则、免责声明和版本灰度。',
    '模型版本、规则版本、命中原因是否对患者展示；内部是否需要可解释性报告。',
    '性能目标、并发量、超时阈值、重试次数、缓存有效期和告警阈值。',
]: add_bullet(doc, x)

# 9 completeness
add_heading(doc, '9 需求完整性校验', 1)
add_table(doc, ['检查项', '结论', '说明'], [
    ('业务目标、用户、范围', '通过', '已明确患者咨询场景、目标用户、三项本次范围和非范围'),
    ('跨页面流程和状态', '通过', '已提供页面地图、总体流程和导诊/查询状态模型'),
    ('字段、校验、权限、脱敏', '存在待确认项', '核心字段已定义，医院接口字段、证件类型和授权方式待确认'),
    ('急诊安全与医学规则', '存在阻塞缺口', '原型提供方向和示例，但生产触发词、组合条件、医学责任人未确认'),
    ('医院接口与数据新鲜度', '存在阻塞缺口', '查询依赖医院系统，接口、错误码、刷新频率和冲突处理未确认'),
    ('异常、重试、重复、过期', '通过', '已覆盖空输入、无法识别、超时、无记录、接口失败、重复提交和会话过期'),
    ('验收路径', '通过', '已覆盖正常、急诊、空输入、接口、权限和示例数据清理'),
    ('原型示例与生产规则隔离', '通过', '已明确示例数据、演示延迟、固定状态和关键词逻辑不能直接生产化'),
], widths=[2.0,1.45,3.3], font_size=8.5, first_col_bold=True)
add_note(doc, '总体结论', '存在待确认项，且存在影响核心流程、医学安全、接口和权限的P0阻塞项。本文件可以作为产品和技术评审基线，但在P0事项确认前，不应直接据此上线或把原型逻辑当作生产医学分诊规则。', fill='FFF2CC', color='7F6000')

add_heading(doc, '附录 A 原型事实清单', 1)
add_para(doc, '以下事实来自用户提供截图及工作目录中的原型页面，用于追踪需求来源，不代表最终生产规则。')
add_table(doc, ['来源', '已观察到的内容', '在需求中的处理'], [
    ('智能导诊.html', '存在欢迎语、文本输入、发送、快速回复、分析中提示、重新导诊和结果对话；原型包含六个导诊方向。', '转化为页面组成、多轮状态和验收样例；医学规则全部标记为需审核'),
    ('智能导诊.html', '原型在无法识别主诉时提示重新描述；分析完成后允许继续补充但不重新计算。', '前者作为当前交互事实；后者标记为生产行为待确认'),
    ('就诊流程指引.html', '包含挂号、候诊、检查、缴费、取药的通用流程说明。', '作为流程指引范围和页面字段依据'),
    ('就诊流程指引.html', '查询表单包含患者姓名、身份证号码、查询按钮、就诊记录、当前进度、报告状态和详情。', '作为信息查询页面基线；身份和接口规则补充安全约束'),
    ('用户截图', '展示患者咨询入口中的流程指引、智能导诊、健康咨询、快速找医生、病情咨询、用药指导和就诊信息查询导航。', '仅将导诊机器人相关的流程指引、智能导诊和信息查询纳入本文件，其他模块不展开'),
], widths=[1.45,3.55,1.75], font_size=8.4, first_col_bold=True)

# Document properties
props = doc.core_properties
props.title = '导诊机器人需求规格说明'
props.subject = '面向患者咨询入口的导诊机器人产品需求'
props.author = 'OpenAI Codex'
props.keywords = '导诊机器人, 智能导诊, 就诊流程, 就诊信息查询, 需求规格'
props.comments = '基于用户提供的导诊机器人原型和截图整理，仅覆盖导诊机器人相关范围。'

# prevent table rows from splitting where possible
for table in doc.tables:
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cantSplit = OxmlElement('w:cantSplit')
        trPr.append(cantSplit)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(str(OUT))
